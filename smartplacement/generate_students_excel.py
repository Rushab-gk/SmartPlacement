import os
import sqlite3
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash

excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database', 'students_master_records.xlsx')
csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database', 'students_master_records.csv')
sqlite_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database', 'smartplacement.sqlite')

students_data = [
    {
        "USN / Roll": "1VA22CS001",
        "Student Name": "Aarav Sharma",
        "Email": "aarav.cs22@college.edu",
        "Branch": "Computer Science & Engineering (CSE)",
        "Grad Year": 2026,
        "CGPA": 8.85,
        "10th %": 92.0,
        "12th %": 90.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Advanced",
        "Technical Skills": "Python, SQL, DSA, Flask",
        "Readiness Score": 88.0
    },
    {
        "USN / Roll": "1VA22CS045",
        "Student Name": "Ananya Rao",
        "Email": "ananya.cs22@college.edu",
        "Branch": "Computer Science & Engineering (CSE)",
        "Grad Year": 2026,
        "CGPA": 7.40,
        "10th %": 85.0,
        "12th %": 82.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Intermediate",
        "Technical Skills": "Java, SQL, HTML/CSS",
        "Readiness Score": 69.0
    },
    {
        "USN / Roll": "1VA22IS012",
        "Student Name": "Rohan Verma",
        "Email": "rohan.is22@college.edu",
        "Branch": "Information Science & Engineering (ISE)",
        "Grad Year": 2026,
        "CGPA": 9.10,
        "10th %": 94.0,
        "12th %": 92.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Advanced",
        "Technical Skills": "Python, Cloud Computing, Docker, DSA",
        "Readiness Score": 92.5
    },
    {
        "USN / Roll": "1VA22IS034",
        "Student Name": "Sneha Patil",
        "Email": "sneha.is22@college.edu",
        "Branch": "Information Science & Engineering (ISE)",
        "Grad Year": 2026,
        "CGPA": 8.20,
        "10th %": 88.0,
        "12th %": 86.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Intermediate",
        "Technical Skills": "Python, SQL, JavaScript",
        "Readiness Score": 77.0
    },
    {
        "USN / Roll": "1VA22EC018",
        "Student Name": "Karthik Nair",
        "Email": "karthik.ec22@college.edu",
        "Branch": "Electronics & Communication (ECE)",
        "Grad Year": 2026,
        "CGPA": 7.80,
        "10th %": 82.0,
        "12th %": 80.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Intermediate",
        "Technical Skills": "C++, Operating Systems, Python",
        "Readiness Score": 73.5
    },
    {
        "USN / Roll": "1VA22EC052",
        "Student Name": "Pooja Hegde",
        "Email": "pooja.ec22@college.edu",
        "Branch": "Electronics & Communication (ECE)",
        "Grad Year": 2026,
        "CGPA": 8.60,
        "10th %": 90.0,
        "12th %": 88.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Advanced",
        "Technical Skills": "C++, Computer Networks, SQL",
        "Readiness Score": 84.0
    },
    {
        "USN / Roll": "1VA22EE009",
        "Student Name": "Aditya Kulkarni",
        "Email": "aditya.ee22@college.edu",
        "Branch": "Electrical & Electronics (EEE)",
        "Grad Year": 2026,
        "CGPA": 7.20,
        "10th %": 78.0,
        "12th %": 75.0,
        "Diploma %": None,
        "Active Backlogs": 1,
        "DSA Level": "Beginner",
        "Technical Skills": "C++, SQL",
        "Readiness Score": 61.0
    },
    {
        "USN / Roll": "1VA22ME025",
        "Student Name": "Vikram Singh",
        "Email": "vikram.me22@college.edu",
        "Branch": "Mechanical Engineering (ME)",
        "Grad Year": 2026,
        "CGPA": 6.90,
        "10th %": 75.0,
        "12th %": 72.0,
        "Diploma %": None,
        "Active Backlogs": 0,
        "DSA Level": "Beginner",
        "Technical Skills": "Python, Operating Systems",
        "Readiness Score": 58.5
    }
]

headers = list(students_data[0].keys())

# 1. Write CSV File
with open(csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=headers)
    writer.writeheader()
    writer.writerows(students_data)

print(f"Generated Student CSV: {csv_path}")

# 2. Write Formatted Excel File (.xlsx)
wb = openpyxl.Workbook()
ws = wb.active
if ws is not None:
    ws.title = "Branchwise Student Master"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for st in students_data:
        row_val = [st[h] for h in headers]
        ws.append(row_val)

    for row in ws.iter_rows(min_row=2, max_row=len(students_data)+1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_idx = col[0].column
        if isinstance(col_idx, int):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(excel_path)
print(f"Generated Student Excel: {excel_path}")

# 3. Seed into SQLite database if DB exists
if os.path.exists(sqlite_path):
    conn = sqlite3.connect(sqlite_path)
    cursor = conn.cursor()
    pwd_hash = generate_password_hash('student123')

    branch_map = {
        "Computer Science & Engineering (CSE)": 1,
        "Information Science & Engineering (ISE)": 2,
        "Electronics & Communication (ECE)": 3,
        "Electrical & Electronics (EEE)": 4,
        "Mechanical Engineering (ME)": 5
    }

    for st in students_data:
        b_id = branch_map.get(st['Branch'], 1)
        try:
            cursor.execute("""
                INSERT OR IGNORE INTO students 
                (usn_or_roll, name, email, password_hash, branch_id, graduation_year, cgpa, tenth_percentage, twelfth_percentage, active_backlogs, dsa_proficiency, projects_count, certifications_count, aptitude_score)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (
                st['USN / Roll'], st['Student Name'], st['Email'], pwd_hash, b_id, st['Grad Year'],
                st['CGPA'], st['10th %'], st['12th %'], st['Active Backlogs'], st['DSA Level'], 2, 1, 80.0
            ))
        except Exception as e:
            print(f"Skipping seeding for {st['USN / Roll']}: {e}")

    conn.commit()
    conn.close()
    print("Seeded student records into SQLite DB.")
