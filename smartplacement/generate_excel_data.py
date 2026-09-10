import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database', 'campus_placement_drives.xlsx')
csv_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'database', 'campus_placement_drives.csv')

data = [
    {
        "Company Name": "TechCorp Systems",
        "Industry": "IT Services & Consulting",
        "Job Role": "Software Engineer Trainee",
        "Package CTC (LPA)": 7.50,
        "Min CGPA": 7.00,
        "Min 10th %": 60.0,
        "Min 12th %": 60.0,
        "Max Backlogs": 0,
        "Grad Year": 2026,
        "Allowed Branches": "CSE, ISE, ECE",
        "Required Skills": "Python, SQL, DSA, HTML/CSS",
        "Drive Date": "2026-10-15",
        "Application Deadline": "2026-10-05",
        "Status": "Upcoming"
    },
    {
        "Company Name": "CloudMatrix Analytics",
        "Industry": "Product / Cloud",
        "Job Role": "Associate Cloud Developer",
        "Package CTC (LPA)": 12.00,
        "Min CGPA": 8.00,
        "Min 10th %": 75.0,
        "Min 12th %": 75.0,
        "Max Backlogs": 0,
        "Grad Year": 2026,
        "Allowed Branches": "CSE, ISE",
        "Required Skills": "Python, DSA, Cloud Computing, Docker",
        "Drive Date": "2026-11-01",
        "Application Deadline": "2026-10-20",
        "Status": "Upcoming"
    },
    {
        "Company Name": "FinTech Solutions Ltd",
        "Industry": "Financial Technology",
        "Job Role": "Backend Developer (Python/Java)",
        "Package CTC (LPA)": 9.00,
        "Min CGPA": 7.50,
        "Min 10th %": 65.0,
        "Min 12th %": 65.0,
        "Max Backlogs": 1,
        "Grad Year": 2026,
        "Allowed Branches": "CSE, ISE, ECE, EEE",
        "Required Skills": "Java, SQL, Flask",
        "Drive Date": "2026-10-25",
        "Application Deadline": "2026-10-12",
        "Status": "Upcoming"
    },
    {
        "Company Name": "CoreElectro Dynamics",
        "Industry": "Embedded Systems",
        "Job Role": "Embedded Systems Engineer",
        "Package CTC (LPA)": 6.50,
        "Min CGPA": 6.50,
        "Min 10th %": 60.0,
        "Min 12th %": 60.0,
        "Max Backlogs": 2,
        "Grad Year": 2026,
        "Allowed Branches": "ECE, EEE, ME",
        "Required Skills": "C++, Operating Systems",
        "Drive Date": "2026-11-10",
        "Application Deadline": "2026-10-30",
        "Status": "Upcoming"
    }
]

headers = list(data[0].keys())

# 1. Generate CSV File
with open(csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=headers)
    writer.writeheader()
    writer.writerows(data)

print(f"Generated CSV file: {csv_path}")

# 2. Generate Formatted Excel (.xlsx) File
wb = openpyxl.Workbook()
ws = wb.active
if ws is not None:
    ws.title = "Campus Placement Drives"

    # Styling definitions
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Write Header Row
    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write Data Rows
    for row_data in data:
        row_val = [row_data[h] for h in headers]
        ws.append(row_val)

    # Style Data Rows & Adjust Column Widths
    for row in ws.iter_rows(min_row=2, max_row=len(data)+1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_idx = col[0].column
        if isinstance(col_idx, int):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save(excel_path)
print(f"Generated Excel file: {excel_path}")
