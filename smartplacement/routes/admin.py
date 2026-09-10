from flask import Blueprint, render_template, request, redirect, url_for, flash, Response, session
from werkzeug.security import generate_password_hash
import csv
import io
from routes.auth import admin_required
from services.db import fetch_one, fetch_all, execute_query

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')

@admin_bp.route('/dashboard')
@admin_required
def dashboard():
    st_res = fetch_one("SELECT COUNT(*) AS count FROM students")
    comp_res = fetch_one("SELECT COUNT(*) AS count FROM companies")
    dr_res = fetch_one("SELECT COUNT(*) AS count FROM placement_drives WHERE status = 'Upcoming'")
    pkg_res = fetch_one("SELECT AVG(package_ctc) AS avg_ctc FROM placement_drives")

    total_students = st_res['count'] if st_res and 'count' in st_res else 0
    total_companies = comp_res['count'] if comp_res and 'count' in comp_res else 0
    active_drives = dr_res['count'] if dr_res and 'count' in dr_res else 0
    raw_avg = pkg_res['avg_ctc'] if pkg_res and 'avg_ctc' in pkg_res and pkg_res['avg_ctc'] is not None else 0.0
    avg_package = float(raw_avg)

    recent_drives = fetch_all("""
        SELECT d.*, c.company_name 
        FROM placement_drives d 
        JOIN companies c ON d.company_id = c.company_id 
        ORDER BY d.created_at DESC LIMIT 5
    """)

    return render_template(
        'admin/dashboard.html',
        total_students=total_students,
        total_companies=total_companies,
        active_drives=active_drives,
        avg_package=round(avg_package, 2),
        recent_drives=recent_drives
    )

@admin_bp.route('/companies', methods=['GET', 'POST'])
@admin_required
def manage_companies():
    if request.method == 'POST':
        company_name = request.form.get('company_name', '').strip()
        industry = request.form.get('industry', '').strip()
        description = request.form.get('description', '').strip()
        website = request.form.get('website', '').strip()
        headquarters = request.form.get('headquarters', '').strip()

        if not company_name or not industry:
            flash('Company name and industry are required.', 'danger')
        else:
            try:
                execute_query(
                    "INSERT INTO companies (company_name, industry, description, website, headquarters) VALUES (%s, %s, %s, %s, %s)",
                    (company_name, industry, description, website, headquarters)
                )
                flash(f'Company "{company_name}" added successfully!', 'success')
            except Exception as e:
                flash(f'Error adding company: {str(e)}', 'danger')

        return redirect(url_for('admin.manage_companies'))

    companies = fetch_all("SELECT * FROM companies ORDER BY company_name ASC")
    return render_template('admin/companies.html', companies=companies)

@admin_bp.route('/companies/delete/<int:company_id>', methods=['POST'])
@admin_required
def delete_company(company_id):
    try:
        execute_query("DELETE FROM companies WHERE company_id = %s", (company_id,))
        flash('Company deleted successfully.', 'success')
    except Exception as e:
        flash(f'Error deleting company: {str(e)}', 'danger')
    return redirect(url_for('admin.manage_companies'))

@admin_bp.route('/drives', methods=['GET', 'POST'])
@admin_required
def manage_drives():
    companies = fetch_all("SELECT * FROM companies ORDER BY company_name ASC")
    branches = fetch_all("SELECT * FROM branches ORDER BY branch_name ASC")
    skills = fetch_all("SELECT * FROM skills ORDER BY skill_name ASC")

    if request.method == 'POST':
        company_id = request.form.get('company_id')
        job_role = request.form.get('job_role', '').strip()
        placement_year = request.form.get('placement_year', 2026)
        package_ctc = request.form.get('package_ctc', 0.0)
        vacancies = request.form.get('vacancies', 0)
        drive_date = request.form.get('drive_date')
        application_deadline = request.form.get('application_deadline')
        
        # Criteria fields
        min_cgpa = request.form.get('minimum_cgpa', 0.0)
        min_tenth = request.form.get('minimum_tenth', 0.0)
        min_twelfth = request.form.get('minimum_twelfth', 0.0)
        max_backlogs = request.form.get('maximum_backlogs', 0)
        other_requirements = request.form.get('other_requirements', '').strip()

        allowed_branch_ids = request.form.getlist('allowed_branches')
        mandatory_skill_ids = request.form.getlist('mandatory_skills')

        if not company_id or not job_role or not drive_date or not application_deadline:
            flash('Please fill in all mandatory drive fields.', 'danger')
            return redirect(url_for('admin.manage_drives'))

        try:
            # 1. Insert Drive
            drive_id = execute_query("""
                INSERT INTO placement_drives 
                (company_id, placement_year, job_role, package_ctc, vacancies, drive_date, application_deadline, status)
                VALUES (%s, %s, %s, %s, %s, %s, %s, 'Upcoming')
            """, (company_id, placement_year, job_role, package_ctc, vacancies, drive_date, application_deadline))

            # 2. Insert Eligibility Criteria
            execute_query("""
                INSERT INTO eligibility_criteria 
                (drive_id, minimum_cgpa, minimum_tenth, minimum_twelfth, maximum_backlogs, graduation_year, other_requirements)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
            """, (drive_id, min_cgpa, min_tenth, min_twelfth, max_backlogs, placement_year, other_requirements))

            # 3. Insert Allowed Branches
            for b_id in allowed_branch_ids:
                execute_query("INSERT INTO drive_branches (drive_id, branch_id) VALUES (%s, %s)", (drive_id, b_id))

            # 4. Insert Required Skills
            for s_id in mandatory_skill_ids:
                execute_query("INSERT INTO drive_skills (drive_id, skill_id, importance) VALUES (%s, %s, 'Mandatory')", (drive_id, s_id))

            # 5. Insert Selection Process Rounds
            round_names = request.form.getlist('round_name[]')
            round_descriptions = request.form.getlist('round_description[]')
            round_idx = 1
            for r_name, r_desc in zip(round_names, round_descriptions):
                if r_name and r_name.strip():
                    execute_query("""
                        INSERT INTO placement_rounds (drive_id, round_number, round_name, description)
                        VALUES (%s, %s, %s, %s)
                    """, (drive_id, round_idx, r_name.strip(), r_desc.strip()))
                    round_idx += 1

            flash('Placement drive created successfully!', 'success')
            return redirect(url_for('admin.manage_drives'))
        except Exception as e:
            flash(f'Error creating placement drive: {str(e)}', 'danger')

    drives = fetch_all("""
        SELECT d.*, c.company_name, ec.minimum_cgpa, ec.maximum_backlogs 
        FROM placement_drives d 
        JOIN companies c ON d.company_id = c.company_id 
        LEFT JOIN eligibility_criteria ec ON d.drive_id = ec.drive_id 
        ORDER BY d.drive_date DESC
    """)

    return render_template('admin/drives.html', drives=drives, companies=companies, branches=branches, skills=skills)

@admin_bp.route('/students/export/<int:drive_id>')
@admin_required
def export_eligible_students(drive_id):
    drive = fetch_one("SELECT d.*, c.company_name FROM placement_drives d JOIN companies c ON d.company_id = c.company_id WHERE d.drive_id = %s", (drive_id,))
    if not drive:
        flash('Drive not found.', 'danger')
        return redirect(url_for('admin.manage_drives'))

    # Fetch eligible students query
    students = fetch_all("""
        SELECT s.usn_or_roll, s.name, s.email, b.branch_name, s.cgpa, s.tenth_percentage, s.twelfth_percentage, s.active_backlogs
        FROM students s
        JOIN branches b ON s.branch_id = b.branch_id
        JOIN drive_branches db ON s.branch_id = db.branch_id
        JOIN eligibility_criteria ec ON ec.drive_id = %s
        WHERE db.drive_id = %s
          AND s.cgpa >= ec.minimum_cgpa
          AND s.tenth_percentage >= ec.minimum_tenth
          AND s.active_backlogs <= ec.maximum_backlogs
          AND s.graduation_year = ec.graduation_year
    """, (drive_id, drive_id))

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['USN/Roll', 'Student Name', 'Email', 'Branch', 'CGPA', '10th %', '12th %', 'Active Backlogs'])

    for st in students:
        writer.writerow([st['usn_or_roll'], st['name'], st['email'], st['branch_name'], st['cgpa'], st['tenth_percentage'], st['twelfth_percentage'], st['active_backlogs']])

    output.seek(0)
    filename = f"Eligible_Students_Drive_{drive_id}_{drive['company_name'].replace(' ', '_')}.csv"
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment;filename={filename}"}
    )
@admin_required
def export_all_drives_csv():
    drives = fetch_all("""
        SELECT d.drive_id, c.company_name, c.industry, d.job_role, d.package_ctc, d.vacancies,
               d.drive_date, d.application_deadline, d.status,
               ec.minimum_cgpa, ec.minimum_tenth, ec.minimum_twelfth, ec.maximum_backlogs, ec.graduation_year
        FROM placement_drives d
        JOIN companies c ON d.company_id = c.company_id
        LEFT JOIN eligibility_criteria ec ON d.drive_id = ec.drive_id
        ORDER BY d.drive_date ASC
    """)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'Drive ID', 'Company Name', 'Industry', 'Job Role', 'Package CTC (LPA)', 'Vacancies',
        'Min CGPA', 'Min 10th %', 'Min 12th %', 'Max Backlogs', 'Grad Year',
        'Drive Date', 'Application Deadline', 'Status'
    ])

    for dr in drives:
        writer.writerow([
            dr['drive_id'], dr['company_name'], dr['industry'], dr['job_role'], dr['package_ctc'], dr['vacancies'],
            dr['minimum_cgpa'], dr['minimum_tenth'], dr['minimum_twelfth'], dr['maximum_backlogs'], dr['graduation_year'],
            dr['drive_date'], dr['application_deadline'], dr['status']
        ])

    output.seek(0)
    filename = "Campus_Placement_Drives_Master_List.csv"
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment;filename={filename}"}
    )

@admin_bp.route('/students/export/all')
@admin_required
def export_all_students_csv():
    students = fetch_all("""
        SELECT s.usn_or_roll, s.name, s.email, b.branch_name, s.graduation_year, s.cgpa,
               s.tenth_percentage, s.twelfth_percentage, s.diploma_percentage, s.active_backlogs, s.dsa_proficiency
        FROM students s
        JOIN branches b ON s.branch_id = b.branch_id
        ORDER BY b.branch_name ASC, s.name ASC
    """)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        'USN/Roll', 'Student Name', 'Email', 'Branch', 'Graduation Year',
        'CGPA', '10th %', '12th %', 'Diploma %', 'Active Backlogs', 'DSA Level'
    ])

    for st in students:
        writer.writerow([
            st['usn_or_roll'], st['name'], st['email'], st['branch_name'], st['graduation_year'],
            st['cgpa'], st['tenth_percentage'], st['twelfth_percentage'], st['diploma_percentage'],
            st['active_backlogs'], st['dsa_proficiency']
        ])

    output.seek(0)
    filename = "Branchwise_Student_Master_Records.csv"
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment;filename={filename}"}
    )

@admin_bp.route('/branches', methods=['GET', 'POST'])
@admin_required
def manage_branches():
    if request.method == 'POST':
        branch_name = request.form.get('branch_name', '').strip()
        branch_code = request.form.get('branch_code', '').strip().upper()

        if not branch_name or not branch_code:
            flash('Branch name and branch code are required.', 'danger')
        else:
            try:
                execute_query(
                    "INSERT INTO branches (branch_name, branch_code) VALUES (%s, %s)",
                    (branch_name, branch_code)
                )
                flash(f'Branch "{branch_name} ({branch_code})" added successfully!', 'success')
            except Exception as e:
                flash(f'Error adding branch: {str(e)}', 'danger')

        return redirect(url_for('admin.manage_branches'))

    branches = fetch_all("""
        SELECT b.*, 
               (SELECT COUNT(*) FROM students s WHERE s.branch_id = b.branch_id) AS student_count,
               (SELECT COUNT(*) FROM drive_branches db WHERE db.branch_id = b.branch_id) AS drive_count
        FROM branches b 
        ORDER BY b.branch_name ASC
    """)
    return render_template('admin/branches.html', branches=branches)

@admin_bp.route('/branches/delete/<int:branch_id>', methods=['POST'])
@admin_required
def delete_branch(branch_id):
    try:
        branch = fetch_one("SELECT * FROM branches WHERE branch_id = %s", (branch_id,))
        if not branch:
            flash('Branch not found.', 'danger')
            return redirect(url_for('admin.manage_branches'))

        st_count = fetch_one("SELECT COUNT(*) AS count FROM students WHERE branch_id = %s", (branch_id,))
        num_students = st_count.get('count', 0) if st_count else 0
        force = request.form.get('force') == 'true'

        if num_students > 0 and not force:
            flash(f'Cannot delete "{branch["branch_name"]}" because {num_students} student(s) are currently enrolled in it.', 'warning')
        else:
            execute_query("DELETE FROM drive_branches WHERE branch_id = %s", (branch_id,))
            if force and num_students > 0:
                execute_query("DELETE FROM student_skills WHERE student_id IN (SELECT student_id FROM students WHERE branch_id = %s)", (branch_id,))
                execute_query("DELETE FROM bookmarks WHERE student_id IN (SELECT student_id FROM students WHERE branch_id = %s)", (branch_id,))
                execute_query("DELETE FROM students WHERE branch_id = %s", (branch_id,))

            execute_query("DELETE FROM branches WHERE branch_id = %s", (branch_id,))
            flash(f'Branch "{branch["branch_name"]} ({branch["branch_code"]})" deleted successfully.', 'success')
    except Exception as e:
        flash(f'Error deleting branch: {str(e)}', 'danger')
    return redirect(url_for('admin.manage_branches'))

@admin_bp.route('/students', methods=['GET', 'POST'])
@admin_required
def manage_students():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        usn_or_roll = request.form.get('usn_or_roll', '').strip().upper()
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', 'student123').strip()
        branch_id = request.form.get('branch_id')
        graduation_year = request.form.get('graduation_year', 2026)
        cgpa = request.form.get('cgpa', 0.0)
        tenth_pct = request.form.get('tenth_percentage', 0.0)
        twelfth_pct = request.form.get('twelfth_percentage') or None
        diploma_pct = request.form.get('diploma_percentage') or None
        active_backlogs = request.form.get('active_backlogs', 0)

        if not name or not usn_or_roll or not email or not branch_id:
            flash('Name, USN/Roll, Email, and Branch are required.', 'danger')
        else:
            try:
                hashed_pw = generate_password_hash(password)
                execute_query("""
                    INSERT INTO students (usn_or_roll, name, email, password_hash, branch_id, graduation_year, cgpa, tenth_percentage, twelfth_percentage, diploma_percentage, active_backlogs)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (usn_or_roll, name, email, hashed_pw, branch_id, graduation_year, cgpa, tenth_pct, twelfth_pct, diploma_pct, active_backlogs))
                flash(f'Student "{name} ({usn_or_roll})" registered successfully!', 'success')
            except Exception as e:
                flash(f'Error registering student: {str(e)}', 'danger')

        return redirect(url_for('admin.manage_students'))

    search_query = request.args.get('search', '').strip()
    selected_branch = request.args.get('branch_id', '').strip()
    selected_year = request.args.get('grad_year', '').strip()

    sql = """
        SELECT s.*, b.branch_name, b.branch_code
        FROM students s
        JOIN branches b ON s.branch_id = b.branch_id
        WHERE 1=1
    """
    params = []

    if search_query:
        sql += " AND (s.name LIKE %s OR s.usn_or_roll LIKE %s OR s.email LIKE %s)"
        params.extend([f"%{search_query}%", f"%{search_query}%", f"%{search_query}%"])
    if selected_branch:
        sql += " AND s.branch_id = %s"
        params.append(selected_branch)
    if selected_year:
        sql += " AND s.graduation_year = %s"
        params.append(selected_year)

    sql += " ORDER BY s.name ASC"

    students = fetch_all(sql, params)
    branches = fetch_all("SELECT * FROM branches ORDER BY branch_name ASC")

    return render_template(
        'admin/students.html',
        students=students,
        branches=branches,
        search_query=search_query,
        selected_branch=selected_branch,
        selected_year=selected_year
    )

@admin_bp.route('/students/edit/<int:student_id>', methods=['POST'])
@admin_required
def edit_student(student_id):
    try:
        name = request.form.get('name', '').strip()
        usn_or_roll = request.form.get('usn_or_roll', '').strip().upper()
        email = request.form.get('email', '').strip().lower()
        branch_id = request.form.get('branch_id')
        graduation_year = request.form.get('graduation_year', 2026)
        cgpa = request.form.get('cgpa', 0.0)
        tenth_pct = request.form.get('tenth_percentage', 0.0)
        twelfth_pct = request.form.get('twelfth_percentage') or None
        diploma_pct = request.form.get('diploma_percentage') or None
        active_backlogs = request.form.get('active_backlogs', 0)

        execute_query("""
            UPDATE students
            SET name = %s, usn_or_roll = %s, email = %s, branch_id = %s, graduation_year = %s,
                cgpa = %s, tenth_percentage = %s, twelfth_percentage = %s, diploma_percentage = %s, active_backlogs = %s
            WHERE student_id = %s
        """, (name, usn_or_roll, email, branch_id, graduation_year, cgpa, tenth_pct, twelfth_pct, diploma_pct, active_backlogs, student_id))
        flash(f'Student profile for "{name}" updated successfully.', 'success')
    except Exception as e:
        flash(f'Error updating student: {str(e)}', 'danger')
    return redirect(url_for('admin.manage_students'))

@admin_bp.route('/students/delete/<int:student_id>', methods=['POST'])
@admin_required
def delete_student(student_id):
    try:
        execute_query("DELETE FROM student_skills WHERE student_id = %s", (student_id,))
        execute_query("DELETE FROM bookmarks WHERE student_id = %s", (student_id,))
        execute_query("DELETE FROM students WHERE student_id = %s", (student_id,))
        flash('Student record deleted successfully.', 'success')
    except Exception as e:
        flash(f'Error deleting student: {str(e)}', 'danger')
    return redirect(url_for('admin.manage_students'))

@admin_bp.route('/profile', methods=['GET', 'POST'])
@admin_required
def profile():
    admin_id = session.get('user_id')
    admin_user = fetch_one("SELECT * FROM admins WHERE admin_id = %s", (admin_id,))
    if not admin_user:
        flash('Admin account not found.', 'danger')
        return redirect(url_for('admin.dashboard'))

    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        email = request.form.get('email', '').strip().lower()
        role = request.form.get('role', 'Placement Director').strip()
        new_password = request.form.get('new_password', '').strip()

        if not name or not email:
            flash('Name and Email are required.', 'danger')
        else:
            try:
                if new_password:
                    hashed_pw = generate_password_hash(new_password)
                    execute_query("""
                        UPDATE admins SET name = %s, email = %s, role = %s, password_hash = %s WHERE admin_id = %s
                    """, (name, email, role, hashed_pw, admin_id))
                else:
                    execute_query("""
                        UPDATE admins SET name = %s, email = %s, role = %s WHERE admin_id = %s
                    """, (name, email, role, admin_id))

                session['name'] = name
                session['email'] = email
                flash('Admin profile updated successfully!', 'success')
                return redirect(url_for('admin.profile'))
            except Exception as e:
                flash(f'Error updating admin profile: {str(e)}', 'danger')

    return render_template('admin/profile.html', admin_user=admin_user)

@admin_bp.route('/students/import', methods=['POST'])
@admin_required
def import_students_excel():
    if 'file' not in request.files:
        flash('No file selected for import.', 'danger')
        return redirect(url_for('admin.manage_students'))

    file = request.files['file']
    if not file or not file.filename:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('admin.manage_students'))

    filename = file.filename.lower()
    branches = fetch_all("SELECT * FROM branches")
    branch_map = {b['branch_code'].upper(): b['branch_id'] for b in branches}
    branch_name_map = {b['branch_name'].lower(): b['branch_id'] for b in branches}

    default_pw_hash = generate_password_hash("student123")
    imported_count = 0
    errors = 0
    rows_data = []

    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode('utf-8-sig', errors='ignore'))
            reader = csv.DictReader(stream)
            for row in reader:
                rows_data.append(row)
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
            sheet = wb.active
            if sheet is not None:
                headers = [str(cell.value or '').strip() for cell in sheet[1]]  # type: ignore[index]
                for row in sheet.iter_rows(min_row=2, values_only=True):       # type: ignore[union-attr]
                    if any(row):
                        row_dict = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
                        rows_data.append(row_dict)

        for row in rows_data:
            key_map = {str(k).strip().lower(): v for k, v in row.items() if k is not None}

            usn = str(key_map.get('usn/roll') or key_map.get('usn') or key_map.get('roll') or key_map.get('usn_or_roll') or '').strip().upper()
            name = str(key_map.get('name') or key_map.get('student name') or key_map.get('student_name') or '').strip()
            email = str(key_map.get('email') or key_map.get('email address') or '').strip().lower()

            if not usn or not name or not email:
                continue

            branch_str = str(key_map.get('branch') or key_map.get('branch_code') or key_map.get('branch code') or 'CSE').strip()
            branch_id = branch_map.get(branch_str.upper()) or branch_name_map.get(branch_str.lower()) or (branches[0]['branch_id'] if branches else 1)

            grad_year = int(key_map.get('graduation year') or key_map.get('graduation_year') or key_map.get('grad year') or 2026)
            cgpa = float(key_map.get('cgpa') or key_map.get('gpa') or 0.0)
            tenth = float(key_map.get('10th %') or key_map.get('tenth_percentage') or key_map.get('10th') or 0.0)
            twelfth_val = key_map.get('12th %') or key_map.get('twelfth_percentage') or key_map.get('12th')
            twelfth = float(twelfth_val) if twelfth_val else None
            diploma_val = key_map.get('diploma %') or key_map.get('diploma_percentage') or key_map.get('diploma')
            diploma = float(diploma_val) if diploma_val else None
            backlogs = int(key_map.get('active backlogs') or key_map.get('active_backlogs') or key_map.get('backlogs') or 0)

            try:
                execute_query("""
                    INSERT INTO students (usn_or_roll, name, email, password_hash, branch_id, graduation_year, cgpa, tenth_percentage, twelfth_percentage, diploma_percentage, active_backlogs)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """, (usn, name, email, default_pw_hash, branch_id, grad_year, cgpa, tenth, twelfth, diploma, backlogs))
                imported_count += 1
            except Exception:
                errors += 1

        flash(f'Successfully imported {imported_count} student records from Excel/CSV! ({errors} skipped/duplicates)', 'success' if imported_count > 0 else 'warning')
    except Exception as e:
        flash(f'Error processing Excel file: {str(e)}', 'danger')

    return redirect(url_for('admin.manage_students'))

@admin_bp.route('/companies/import', methods=['POST'])
@admin_required
def import_companies_excel():
    if 'file' not in request.files:
        flash('No file selected for import.', 'danger')
        return redirect(url_for('admin.manage_companies'))

    file = request.files['file']
    if not file or not file.filename:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('admin.manage_companies'))

    filename = file.filename.lower()
    imported_count = 0
    errors = 0
    rows_data = []

    try:
        if filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode('utf-8-sig', errors='ignore'))
            reader = csv.DictReader(stream)
            for row in reader:
                rows_data.append(row)
        elif filename.endswith('.xlsx') or filename.endswith('.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
            sheet = wb.active
            if sheet is not None:
                headers = [str(cell.value or '').strip() for cell in sheet[1]]  # type: ignore[index]
                for row in sheet.iter_rows(min_row=2, values_only=True):       # type: ignore[union-attr]
                    if any(row):
                        row_dict = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
                        rows_data.append(row_dict)

        for row in rows_data:
            key_map = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
            c_name = str(key_map.get('company name') or key_map.get('company_name') or key_map.get('company') or '').strip()
            industry = str(key_map.get('industry') or key_map.get('sector') or 'IT Services').strip()
            description = str(key_map.get('description') or key_map.get('about') or '').strip()
            website = str(key_map.get('website') or key_map.get('url') or '').strip()
            headquarters = str(key_map.get('headquarters') or key_map.get('location') or '').strip()

            if not c_name:
                continue

            try:
                execute_query("""
                    INSERT INTO companies (company_name, industry, description, website, headquarters)
                    VALUES (%s, %s, %s, %s, %s)
                """, (c_name, industry, description, website, headquarters))
                imported_count += 1
            except Exception:
                errors += 1

        flash(f'Successfully imported {imported_count} companies from Excel/CSV! ({errors} skipped/duplicates)', 'success' if imported_count > 0 else 'warning')
    except Exception as e:
        flash(f'Error processing Excel file: {str(e)}', 'danger')

    return redirect(url_for('admin.manage_companies'))

@admin_bp.route('/drives/delete/<int:drive_id>', methods=['POST'])
@admin_required
def delete_drive(drive_id):
    try:
        execute_query("DELETE FROM eligibility_criteria WHERE drive_id = %s", (drive_id,))
        execute_query("DELETE FROM drive_branches WHERE drive_id = %s", (drive_id,))
        execute_query("DELETE FROM drive_skills WHERE drive_id = %s", (drive_id,))
        execute_query("DELETE FROM placement_rounds WHERE drive_id = %s", (drive_id,))
        execute_query("DELETE FROM bookmarks WHERE drive_id = %s", (drive_id,))
        execute_query("DELETE FROM placement_drives WHERE drive_id = %s", (drive_id,))
        flash('Placement drive deleted successfully.', 'success')
    except Exception as e:
        flash(f'Error deleting placement drive: {str(e)}', 'danger')
    return redirect(url_for('admin.manage_drives'))

@admin_bp.route('/skills/add', methods=['POST'])
@admin_required
def add_skill():
    skill_name = request.form.get('skill_name', '').strip()
    skill_category = request.form.get('skill_category', 'Technical').strip()

    if not skill_name:
        flash('Skill name is required.', 'danger')
    else:
        try:
            execute_query("INSERT INTO skills (skill_name, skill_category) VALUES (%s, %s)", (skill_name, skill_category))
            flash(f'Skill "{skill_name}" added successfully!', 'success')
        except Exception as e:
            flash(f'Error adding skill: {str(e)}', 'danger')

    return redirect(request.referrer or url_for('admin.manage_drives'))

